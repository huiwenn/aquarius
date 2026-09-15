#!/usr/bin/env python3
"""
POP909 Dataset Inspector
========================
Walks the POP909 dataset directory, characterizes all files,
extracts metadata from index.xlsx, analyzes MIDI files and
annotation files, and reports statistics.

Usage:
    conda activate py312
    python src/inspectors/pop909_inspect.py
"""

import os
import sys
import json
from pathlib import Path
from collections import Counter, defaultdict

import openpyxl
import pretty_midi

# ── Configuration ──────────────────────────────────────────────────
DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "raw" / "pop909"
POP909_DIR = DATA_DIR / "POP909"

# ── Helpers ────────────────────────────────────────────────────────

def parse_beat_file(filepath):
    """Parse beat annotation file (beat_audio.txt or beat_midi.txt)."""
    beats = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            parts = line.strip().split()
            if len(parts) >= 2:
                beats.append({
                    "time": float(parts[0]),
                    "beat_position": float(parts[1]),
                })
            elif len(parts) == 3:
                beats.append({
                    "time": float(parts[0]),
                    "beat_position": float(parts[1]),
                    "downbeat": float(parts[2]),
                })
    return beats


def parse_chord_file(filepath):
    """Parse chord annotation file."""
    chords = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            parts = line.strip().split()
            if len(parts) >= 3:
                chords.append({
                    "start": float(parts[0]),
                    "end": float(parts[1]),
                    "chord": parts[2],
                })
    return chords


def parse_key_file(filepath):
    """Parse key annotation file."""
    keys = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            parts = line.strip().split()
            if len(parts) >= 3:
                keys.append({
                    "start": float(parts[0]),
                    "end": float(parts[1]),
                    "key": parts[2],
                })
    return keys


def analyze_midi(filepath):
    """Analyze a MIDI file and return track/instrument info."""
    try:
        midi = pretty_midi.PrettyMIDI(str(filepath))
        tracks = []
        for inst in midi.instruments:
            tracks.append({
                "name": inst.name,
                "program": inst.program,
                "is_drum": inst.is_drum,
                "num_notes": len(inst.notes),
            })
        duration = midi.get_end_time()
        tempos = midi.get_tempo_changes()
        tempo_values = list(tempos[1]) if len(tempos[1]) > 0 else []
        return {
            "duration_sec": duration,
            "num_tracks": len(tracks),
            "tracks": tracks,
            "tempo_values": tempo_values,
            "time_signatures": [
                {"numerator": ts.numerator, "denominator": ts.denominator}
                for ts in midi.time_signature_changes
            ],
            "key_signatures": [
                {"key_number": ks.key_number, "key_name": pretty_midi.key_number_to_key_name(ks.key_number)}
                for ks in midi.key_signature_changes
            ],
        }
    except Exception as e:
        return {"error": str(e)}


def inspect_index_xlsx(filepath):
    """Read and analyze the index.xlsx metadata file."""
    wb = openpyxl.load_workbook(str(filepath))
    ws = wb.active

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))

    # Analyze columns
    col_stats = {}
    for col in headers:
        values = [r[col] for r in rows]
        non_null = [v for v in values if v is not None]
        col_stats[col] = {
            "total": len(values),
            "non_null": len(non_null),
            "null": len(values) - len(non_null),
            "unique": len(set(non_null)),
            "examples": list(set(non_null))[:5],
        }

    return headers, rows, col_stats


# ── Main Inspection ───────────────────────────────────────────────

def main():
    print("=" * 70)
    print("POP909 Dataset Inspection Report")
    print("=" * 70)

    # ── 1. Directory structure ──
    print("\n## 1. Directory Structure")

    all_files = []
    file_types = Counter()
    for root, dirs, files in os.walk(DATA_DIR):
        for f in files:
            fp = Path(root) / f
            ext = fp.suffix.lower()
            size = fp.stat().st_size
            rel = fp.relative_to(DATA_DIR)
            all_files.append({"path": str(rel), "ext": ext, "size": size})
            file_types[ext] += 1

    print(f"Total files: {len(all_files)}")
    total_size = sum(f["size"] for f in all_files)
    print(f"Total size: {total_size / 1024 / 1024:.1f} MB")
    print(f"\nFile type distribution:")
    for ext, count in file_types.most_common():
        sizes = [f["size"] for f in all_files if f["ext"] == ext]
        print(f"  {ext or '(no ext)'}: {count} files, "
              f"total {sum(sizes)/1024/1024:.2f} MB, "
              f"avg {sum(sizes)/len(sizes)/1024:.1f} KB")

    # ── 2. Song directory structure ──
    print("\n## 2. Song Directory Structure")

    song_dirs = sorted([
        d for d in POP909_DIR.iterdir()
        if d.is_dir() and d.name.isdigit()
    ], key=lambda x: int(x.name))

    print(f"Number of song directories: {len(song_dirs)}")
    print(f"Song ID range: {song_dirs[0].name} - {song_dirs[-1].name}")

    # Check consistency of file structure across songs
    file_patterns = Counter()
    version_counts = Counter()
    missing_files = defaultdict(list)
    expected_files = {"beat_audio.txt", "beat_midi.txt", "chord_audio.txt",
                      "chord_midi.txt", "key_audio.txt"}

    for sd in song_dirs:
        sid = sd.name
        files_in_dir = set(f.name for f in sd.iterdir() if f.is_file())
        pattern = frozenset(files_in_dir)
        file_patterns[pattern] += 1

        for ef in expected_files:
            if ef not in files_in_dir:
                missing_files[ef].append(sid)

        versions_dir = sd / "versions"
        if versions_dir.exists():
            v_count = len(list(versions_dir.glob("*.mid")))
            version_counts[v_count] += 1
        else:
            version_counts[0] += 1

    print(f"\nUnique file patterns across songs: {len(file_patterns)}")
    for pattern, count in file_patterns.most_common(5):
        print(f"  {count} songs have: {sorted(pattern)}")

    print(f"\nMissing annotation files:")
    for ef, sids in missing_files.items():
        if sids:
            print(f"  {ef}: missing in {len(sids)} songs (e.g., {sids[:5]})")
        else:
            print(f"  {ef}: present in all songs")

    print(f"\nVersion MIDI counts per song:")
    for vc, count in sorted(version_counts.items()):
        print(f"  {vc} versions: {count} songs")

    # ── 3. Index.xlsx Analysis ──
    print("\n## 3. Index.xlsx Metadata Analysis")

    index_path = DATA_DIR / "index.xlsx"
    if not index_path.exists():
        index_path = POP909_DIR / "index.xlsx"

    headers, rows, col_stats = inspect_index_xlsx(index_path)
    print(f"Columns: {headers}")
    print(f"Total rows: {len(rows)}")

    for col, stats in col_stats.items():
        print(f"\n  Column: '{col}'")
        print(f"    Non-null: {stats['non_null']}/{stats['total']}")
        print(f"    Unique values: {stats['unique']}")
        examples = stats['examples']
        if len(examples) <= 10:
            print(f"    Values: {sorted(examples, key=str)}")
        else:
            print(f"    Examples: {examples[:5]}")

    # Artist distribution
    artist_counts = Counter(r["artist"] for r in rows if r.get("artist"))
    print(f"\n  Artist statistics:")
    print(f"    Total unique artists: {len(artist_counts)}")
    print(f"    Most common artists:")
    for artist, count in artist_counts.most_common(15):
        print(f"      {artist}: {count} songs")

    # modify_times distribution
    modify_counts = Counter(str(r["modify_times"]) for r in rows if r.get("modify_times"))
    print(f"\n  modify_times distribution:")
    for mt, count in sorted(modify_counts.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 0):
        print(f"    {mt} modifications: {count} songs")

    # Beats per measure and quavers
    bpm_counts = Counter(str(r["num_beats_per_measure"]) for r in rows)
    qpb_counts = Counter(str(r["num_quavers_per_beat"]) for r in rows)
    print(f"\n  num_beats_per_measure: {dict(bpm_counts)}")
    print(f"  num_quavers_per_beat: {dict(qpb_counts)}")

    # ── 4. MIDI Analysis (sample) ──
    print("\n## 4. MIDI File Analysis")

    # Analyze main MIDI files (sample of 50 evenly spaced)
    sample_indices = list(range(0, len(song_dirs), max(1, len(song_dirs) // 50)))[:50]

    midi_stats = {
        "durations": [],
        "track_counts": Counter(),
        "track_names": Counter(),
        "tempo_values": [],
        "time_sigs": Counter(),
        "key_sigs": Counter(),
        "notes_per_track": defaultdict(list),
        "errors": [],
    }

    for idx in sample_indices:
        sd = song_dirs[idx]
        sid = sd.name
        mid_file = sd / f"{sid}.mid"
        if mid_file.exists():
            info = analyze_midi(mid_file)
            if "error" in info:
                midi_stats["errors"].append((sid, info["error"]))
                continue
            midi_stats["durations"].append(info["duration_sec"])
            midi_stats["track_counts"][info["num_tracks"]] += 1
            for t in info["tracks"]:
                midi_stats["track_names"][t["name"]] += 1
                midi_stats["notes_per_track"][t["name"]].append(t["num_notes"])
            if info["tempo_values"]:
                midi_stats["tempo_values"].extend(info["tempo_values"])
            for ts in info["time_signatures"]:
                midi_stats["time_sigs"][f"{ts['numerator']}/{ts['denominator']}"] += 1
            for ks in info["key_signatures"]:
                midi_stats["key_sigs"][ks["key_name"]] += 1

    durations = midi_stats["durations"]
    if durations:
        print(f"  Sample size: {len(durations)} main MIDI files")
        print(f"  Duration range: {min(durations):.1f}s - {max(durations):.1f}s")
        print(f"  Mean duration: {sum(durations)/len(durations):.1f}s")
        print(f"  Median duration: {sorted(durations)[len(durations)//2]:.1f}s")

    print(f"\n  Track count distribution: {dict(midi_stats['track_counts'])}")
    print(f"  Track names: {dict(midi_stats['track_names'])}")

    if midi_stats["tempo_values"]:
        tempos = midi_stats["tempo_values"]
        print(f"\n  Tempo range: {min(tempos):.1f} - {max(tempos):.1f} BPM")
        print(f"  Mean tempo: {sum(tempos)/len(tempos):.1f} BPM")

    print(f"\n  Time signatures: {dict(midi_stats['time_sigs'])}")
    print(f"  MIDI key signatures: {dict(midi_stats['key_sigs'])}")

    for tname, notes in midi_stats["notes_per_track"].items():
        print(f"  Notes in '{tname}': min={min(notes)}, max={max(notes)}, avg={sum(notes)/len(notes):.0f}")

    if midi_stats["errors"]:
        print(f"\n  MIDI parse errors: {len(midi_stats['errors'])}")
        for sid, err in midi_stats["errors"][:5]:
            print(f"    Song {sid}: {err}")

    # ── 5. Version MIDI Analysis ──
    print("\n## 5. Version MIDI Analysis (sample)")

    version_track_names = Counter()
    version_note_counts = []
    sample_count = 0
    for idx in sample_indices[:20]:
        sd = song_dirs[idx]
        versions_dir = sd / "versions"
        if versions_dir.exists():
            for vf in versions_dir.glob("*.mid"):
                info = analyze_midi(vf)
                if "error" not in info:
                    sample_count += 1
                    for t in info["tracks"]:
                        version_track_names[t["name"]] += 1
                        version_note_counts.append(t["num_notes"])

    print(f"  Sampled {sample_count} version MIDI files")
    print(f"  Track names in versions: {dict(version_track_names)}")

    # ── 6. Annotation File Analysis ──
    print("\n## 6. Annotation File Analysis")

    # Chord vocabulary
    all_chords = Counter()
    chord_file_count = 0
    for sd in song_dirs:
        for cf_name in ["chord_audio.txt", "chord_midi.txt"]:
            cf = sd / cf_name
            if cf.exists():
                chords = parse_chord_file(cf)
                chord_file_count += 1
                for c in chords:
                    all_chords[c["chord"]] += 1

    print(f"  Chord annotation files analyzed: {chord_file_count}")
    print(f"  Unique chord labels: {len(all_chords)}")
    print(f"  Most common chords:")
    for chord, count in all_chords.most_common(20):
        print(f"    {chord}: {count}")

    # Chord quality types
    chord_qualities = Counter()
    for chord_label in all_chords:
        if chord_label == "N":
            chord_qualities["N (no chord)"] += all_chords[chord_label]
        elif ":" in chord_label:
            quality = chord_label.split(":")[1]
            chord_qualities[quality] += all_chords[chord_label]
        else:
            chord_qualities["major (implied)"] += all_chords[chord_label]

    print(f"\n  Chord quality distribution:")
    for quality, count in chord_qualities.most_common():
        print(f"    {quality}: {count}")

    # Key annotations
    all_keys = Counter()
    key_change_counts = Counter()
    for sd in song_dirs:
        kf = sd / "key_audio.txt"
        if kf.exists():
            keys = parse_key_file(kf)
            key_change_counts[len(keys)] += 1
            for k in keys:
                all_keys[k["key"]] += 1

    print(f"\n  Key labels found: {len(all_keys)}")
    print(f"  Key distribution:")
    for key, count in all_keys.most_common():
        print(f"    {key}: {count}")

    print(f"\n  Key changes per song:")
    for num_changes, count in sorted(key_change_counts.items()):
        print(f"    {num_changes} key(s): {count} songs")

    # Beat annotation analysis
    beat_positions = Counter()
    beat_file_formats = Counter()
    for sd in song_dirs[:50]:  # Sample
        for bf_name in ["beat_audio.txt", "beat_midi.txt"]:
            bf = sd / bf_name
            if bf.exists():
                with open(bf, "r") as f:
                    first_line = f.readline().strip()
                    num_cols = len(first_line.split())
                    beat_file_formats[f"{bf_name}: {num_cols} columns"] += 1
                beats = parse_beat_file(bf)
                for b in beats:
                    beat_positions[b["beat_position"]] += 1

    print(f"\n  Beat file formats (sample of 50 songs):")
    for fmt, count in beat_file_formats.most_common():
        print(f"    {fmt}: {count} files")

    print(f"  Beat positions found: {sorted(set(beat_positions.keys()))}")

    # ── 7. Summary Statistics ──
    print("\n## 7. Summary Statistics")
    print(f"  Total songs: {len(song_dirs)}")
    print(f"  Total unique artists: {len(artist_counts)}")
    total_versions = sum(
        len(list((sd / "versions").glob("*.mid")))
        for sd in song_dirs if (sd / "versions").exists()
    )
    print(f"  Total version MIDI files: {total_versions}")
    print(f"  Total main MIDI files: {sum(1 for sd in song_dirs if (sd / f'{sd.name}.mid').exists())}")
    print(f"  Total annotation files (txt): {file_types.get('.txt', 0)}")
    print(f"  Dataset size on disk: {total_size / 1024 / 1024:.1f} MB")

    print("\n" + "=" * 70)
    print("Inspection complete.")


if __name__ == "__main__":
    main()
